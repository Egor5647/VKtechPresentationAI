#!/usr/bin/env python3
"""Small live API check using synthetic data only; consumes provider credits."""
import argparse
import io
import json
import os
import tempfile
from pathlib import Path

from PIL import Image
from pydantic import BaseModel

from vktech.model import ModelGateway, ModelUnavailable


class Probe(BaseModel):
    answer: str


def check_pipeline(gateway, folder, image_path):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from vktech.contracts import Claim, ContentIR, ImageSelection, PresentationPlan
    from vktech.planning import regenerate_slide_with_model
    from vktech.settings import artifact_path
    from vktech.store import Store
    from vktech.template import import_template
    from vktech.worker import execute

    os.environ['DATA_DIR']=str(folder/'artifacts')
    os.environ['DATABASE_URL']='sqlite:///'+str(folder/'smoke.sqlite')
    presentation=Presentation();presentation.slide_width=Inches(13.333333);presentation.slide_height=Inches(7.5)
    for _ in range(3):
        slide=presentation.slides.add_slide(presentation.slide_layouts[6])
        for x,y,w,h,text,size in ((.7,.4,12,1,'Заголовок',28),(.7,1.9,11.7,4.8,'Текст шаблона',18)):
            box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));box.text=text
            font=box.text_frame.paragraphs[0].font;font.name='DejaVu Sans';font.size=Pt(size);font.color.rgb=RGBColor.from_string('202020')
    buffer=io.BytesIO();presentation.save(buffer);raw=buffer.getvalue();design=import_template(raw)
    content=ContentIR(id='synthetic-api-check',title='Подготовка презентации',claims=[
        Claim(id='c1',text='Шаблон PPTX задаёт фирменные цвета и шрифты презентации.',source='synthetic test'),
        Claim(id='c2',text='Планировщик связывает тезисы с исходными материалами, затем верстальщик создаёт редактируемые слайды.',source='synthetic test'),
        Claim(id='c3',text='Пользователь выбирает компоновки и скачивает презентацию в PPTX, PDF или HTML.',source='synthetic test'),
    ])
    store=Store()
    try:
        tid=store.save_record('template','synthetic.pptx',raw,design.model_dump(),'pptx')
        cid=store.save_record('content','synthetic.json',content.model_dump_json().encode(),content.model_dump(),'json')
        jid=store.enqueue('generate',{'template_id':tid,'content_id':cid,'brief':'Объясни последовательность подготовки презентации на основе всех исходных фактов.','slide_count':3,'generate_images':False})
        execute(store,store.claim(),gateway)
        job=store.job(jid)
        if job.state!='ready':raise ValueError(f'Pipeline check: {job.state}: {job.error}')
        result=json.loads(job.result)
        if not all(artifact_path(result['presentation'][fmt]).is_file() for fmt in ('pptx','pdf','html')):
            raise ValueError('Pipeline did not produce all export formats')
        print(f'pipeline: OK (3 slides, 3 variants, PPTX/PDF/HTML, {result["elapsed_seconds"]} s)',flush=True)
        plan=PresentationPlan.model_validate_json(artifact_path(result['artifacts']['plan']).read_bytes())
        regenerate_slide_with_model(gateway,plan.slides[1],content,'Сделай формулировки яснее, сохрани факты.')
        print('regenerate_slide: OK (SlideRevision contract)',flush=True)
        gateway.structured('image_selection',{'slide_id':'synthetic','title':'Синий цвет','message':'Синяя карточка на тестовом изображении.','visual_contract':{'goal':'Показать синий цвет','entities':['синяя карточка'],'relations':[],'forbidden':[]},'candidate_indices':[0]},ImageSelection,[image_path])
        print('image_selection: OK (ImageSelection contract)',flush=True)
    finally:
        store.engine.dispose()


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--image', action='store_true', help='Also generate one Z-Image illustration (paid API call)')
    parser.add_argument('--pipeline', action='store_true', help='Use production schemas and build a three-slide synthetic deck; needs LibreOffice and Poppler')
    args=parser.parse_args()
    gateway=ModelGateway();gateway.cache_enabled=False
    try:
        with tempfile.TemporaryDirectory(prefix='vktech-api-check-') as folder:
            path=Path(folder)/'blue.png';Image.new('RGB',(128,72),'blue').save(path)
            if args.pipeline:
                check_pipeline(gateway,Path(folder),path)
            else:
                for role in ('planning','regenerate_slide','vision_audit','image_selection'):
                    result=gateway.structured(role,{'check':'Connection test using synthetic data. Return answer="ok" in the supplied JSON schema. The attached test image is blue.'},Probe,[path])
                    if not result.answer.strip():raise ValueError('Empty API check answer')
                    print(f'{role}: OK ({gateway.model})', flush=True)
            if args.image:
                gateway.image('A simple blue geometric cube on a pale background, no text, wide editorial illustration.',Path(folder)/'generated.png')
                print(f'text_to_image: OK ({gateway.image_endpoint.model})', flush=True)
    except (ModelUnavailable,ValueError) as exc:
        print(str(exc))
        raise SystemExit(1) from None
    finally:
        print(json.dumps({'calls':gateway.calls},ensure_ascii=False,indent=2))
        gateway.client.close()


if __name__=='__main__':main()
