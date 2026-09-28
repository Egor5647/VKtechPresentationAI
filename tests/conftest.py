from __future__ import annotations
import io
import os
import pytest
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from vktech.contracts import ContentIR,PresentationPlan,PlanSlide
from vktech.settings import ROOT


@pytest.fixture(autouse=True)
def isolated_model_environment(monkeypatch, tmp_path):
    """Keep local API credentials and production data out of offline tests."""
    for name in list(os.environ):
        if name.startswith(('MODEL_', 'POLZA_', 'T2I_', 'VK_', 'LOCAL_MODEL_')):
            monkeypatch.delenv(name)
    monkeypatch.setenv('DATA_DIR', str(tmp_path / 'data'))
    monkeypatch.setenv('DATABASE_URL', 'sqlite:///' + str(tmp_path / 'service.sqlite'))


@pytest.fixture
def template_bytes():
    prs=Presentation();prs.slide_width=Inches(10);prs.slide_height=Inches(5.625)
    for i in range(6):
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor(255,255,255)
        title=slide.shapes.add_textbox(Inches(.5),Inches(.3),Inches(9),Inches(.8));title.text='Заголовок'
        p=title.text_frame.paragraphs[0];p.font.name='Arial';p.font.size=Pt(28);p.font.color.rgb=RGBColor.from_string('202020')
        body=slide.shapes.add_textbox(Inches(.5+i*.04),Inches(1.4),Inches(8.8-i*.08),Inches(3.6));body.text='Текст шаблона'
        p=body.text_frame.paragraphs[0];p.font.name='Arial';p.font.size=Pt(18);p.font.color.rgb=RGBColor.from_string('202020')
    b=io.BytesIO();prs.save(b);return b.getvalue()


@pytest.fixture
def content():return ContentIR.model_validate_json((ROOT/'fixtures/content.example.json').read_bytes())


@pytest.fixture
def empty_placeholder_template():
    prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
    layout=prs.slide_layouts[1]
    title=layout.placeholders[0]
    title.left=Inches(.5);title.top=Inches(.4);title.width=Inches(9);title.height=Pt(30)
    title.text_frame.paragraphs[0].font.name='DejaVu Sans'
    title.text_frame.paragraphs[0].font.size=Pt(20)
    body=layout.placeholders[1]
    body.text_frame.paragraphs[0].font.name='DejaVu Sans'
    body.text_frame.paragraphs[0].font.size=Pt(18)
    prs.slides.add_slide(layout)
    out=io.BytesIO();prs.save(out);return out.getvalue()


@pytest.fixture
def plan(content):
    return PresentationPlan(slides=[PlanSlide(id=f'slide-{i+1}',title=c.text,message=c.text,claim_ids=[c.id],dataset_id='metrics' if i==5 else None,visual='chart' if i==5 else 'sequence' if i==3 else 'none',role='cover' if i==0 else 'content') for i,c in enumerate(content.claims)])


class FixtureGateway:
    """Explicit test double. Never reachable from the production API/config."""
    profile='test'
    manifest={'test_double':True}
    calls=[]
    def __init__(self,plan):self.plan=plan
    def structured(self,role,payload,schema,images=None):
        if role=='planning':return self.plan
        from vktech.contracts import ContextualReport
        return ContextualReport(issues=[])
