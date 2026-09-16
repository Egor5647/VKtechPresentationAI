from __future__ import annotations

import io

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from vktech.content import import_content


def test_pptx_content_groups_slides_notes_and_embedded_assets(tmp_path):
    prs=Presentation()
    first=prs.slides.add_slide(prs.slide_layouts[6])
    box=first.shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(2))
    box.text='Определение массива'
    box.text_frame.add_paragraph().text='Элементы имеют один тип'
    first.notes_slide.notes_text_frame.text='Методическая заметка'

    second=prs.slides.add_slide(prs.slide_layouts[6])
    image=Image.new('RGB',(40,30),'red');raw=io.BytesIO();image.save(raw,'PNG')
    second.shapes.add_picture(io.BytesIO(raw.getvalue()),Inches(1),Inches(1),Inches(4),Inches(3))
    deck=io.BytesIO();prs.save(deck)

    saved=[]
    def writer(data,name,extension):
        path=tmp_path/name;path.write_bytes(data);saved.append((path,extension));return str(path)

    content=import_content(deck.getvalue(),'lesson.pptx',writer)
    assert [c.id for c in content.claims]==['claim-slide-1','claim-slide-1-notes','claim-slide-2']
    assert content.claims[0].text=='Определение массива\nЭлементы имеют один тип'
    assert content.claims[0].required is True
    assert content.claims[1].text=='Методическая заметка' and content.claims[1].required is False
    assert content.claims[2].required is False and 'без извлекаемого текста' in content.claims[2].text
    assert len(content.assets)==1 and len(saved)==1
    assert content.assets[0].claim_ids==['claim-slide-2']
    assert content.assets[0].media_type=='image/png'
    assert saved[0][0].read_bytes()==raw.getvalue()
