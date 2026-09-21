from __future__ import annotations
import base64
import copy
import html
import io
import json
import os
import posixpath
import re
import subprocess
import tempfile
from pathlib import Path
from lxml import etree as E
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from .contracts import SceneIR, DesignIR
from .graph_layout import semantic_graph,layout_graph
from .opc import Package, NS, serialize, relpath
from .smartart import inject
from .settings import artifact_path


_SUBSCRIPT=str.maketrans({'0':'₀','1':'₁','2':'₂','3':'₃','4':'₄','5':'₅','6':'₆','7':'₇','8':'₈','9':'₉','a':'ₐ','e':'ₑ','h':'ₕ','i':'ᵢ','j':'ⱼ','k':'ₖ','l':'ₗ','m':'ₘ','n':'ₙ','o':'ₒ','p':'ₚ','r':'ᵣ','s':'ₛ','t':'ₜ','u':'ᵤ','v':'ᵥ','x':'ₓ'})


def _math_display(value: str) -> str:
    """Use readable Unicode subscripts in rendered formulas, keeping IR stable."""
    def replace(match):
        suffix=match.group(1)
        converted=suffix.casefold().translate(_SUBSCRIPT)
        return converted if len(converted)==len(suffix) else '_'+suffix
    return re.sub(r'_([A-Za-z0-9]+)',replace,value)


def _card_font_size(node,w,h,text,minimum=12):
    """Fit a complete card label without clipping or shortening its text."""
    from .audit import font_for
    from PIL import ImageFont
    width=max(1,w/914400*96-24);height=max(1,h/914400*96-16)
    maximum=float(node.style.size)
    candidates=[];size=maximum
    while size>=minimum:
        candidates.append(size);size-=1
    for size in candidates:
        font=font_for(node.style.font,max(1,round(size*96/72))) or ImageFont.load_default(size=max(1,round(size*96/72)))
        if any(font.getlength(word)>width for paragraph in text.split('\n') for word in paragraph.split()):continue
        lines=0
        for paragraph in text.split('\n'):
            current='';lines+=1
            for word in paragraph.split():
                candidate=(current+' '+word).strip()
                if current and font.getlength(candidate)>width:lines+=1;current=word
                else:current=candidate
        if lines*size*96/72*1.16<=height:return size
    return float(minimum)


def _card(out,node,x,y,w,h,text,name,fill,foreground,bold=False):
    text=_math_display(text)
    sh=out.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,x,y,w,h)
    sh.name=name;sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill)
    sh.line.color.rgb=RGBColor.from_string(node.data.get('accent','0077FF'));sh.line.width=Pt(1.4)
    _plain_shape(sh)
    tf=sh.text_frame;tf.clear();tf.word_wrap=not (len(text)<=14 and '\n' not in text);tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    tf.margin_left=tf.margin_right=Pt(9);tf.margin_top=tf.margin_bottom=Pt(6)
    p=tf.paragraphs[0];p.text=text;p.alignment=PP_ALIGN.CENTER
    p.font.name=node.style.font;p.font.size=Pt(_card_font_size(node,w,h,text));p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(foreground)
    return sh


def _plain_shape(shape):
    """Suppress theme effects so graph objects render flat in Office and LO."""
    shape.shadow.inherit=False
    sppr=shape._element.spPr
    for child in list(sppr):
        if child.tag.endswith('effectLst') or child.tag.endswith('effectDag'):sppr.remove(child)
    sppr.append(OxmlElement('a:effectLst'))
    style=shape._element.find('p:style',shape._element.nsmap)
    if style is not None:
        effect=style.find('a:effectRef',shape._element.nsmap)
        if effect is not None:effect.set('idx','0')


def _arrow(line):
    """Add a restrained native PowerPoint arrowhead to a connector."""
    ln=line._get_or_add_ln()
    ln.set('cap','rnd')
    for existing in list(ln):
        if existing.tag.endswith('tailEnd'):ln.remove(existing)
    arrow=OxmlElement('a:tailEnd');arrow.set('type','stealth');arrow.set('w','sm');arrow.set('len','med');ln.append(arrow)


def _graph_node_bounds(placed,gx,gy,gw,gh):
    """Project normalized graph geometry while preserving true circles."""
    nx=gx+round(placed.x*gw);ny=gy+round(placed.y*gh);nw=round(placed.w*gw);nh=round(placed.h*gh)
    if placed.shape=='circle':
        diameter=min(nw,nh);nx+=(nw-diameter)//2;ny+=(nh-diameter)//2;nw=nh=diameter
    return nx,ny,nw,nh


def _graph_boundary_point(bounds,placed,toward_x,toward_y):
    """Intersect a ray from a node centre with its rendered contour."""
    nx,ny,nw,nh=bounds[placed.id];cx,cy=nx+nw/2,ny+nh/2
    dx=toward_x-cx;dy=toward_y-cy
    if not dx and not dy:return cx,cy
    if placed.shape=='circle':scale=(nw/2)/max(1,(dx*dx+dy*dy)**.5)
    else:scale=1/max(abs(dx)/(nw/2),abs(dy)/(nh/2))
    return cx+dx*scale,cy+dy*scale


def _render_semantic_graph(out,node,xywh,layout_name,items,accent,surface,dark,white):
    spec=semantic_graph(layout_name,items,node.data.get('graph'))
    if spec is None:return False
    result=layout_graph(spec)
    if result.crossings>spec.rules.max_crossings:raise ValueError(f'Graph layout has {result.crossings} edge crossings')
    x,y,w,h=xywh
    caption_band=round(w*.19) if result.layer_captions and result.direction in {'TB','BT'} else 0
    footer_h=round(h*.11) if result.footer else 0
    legend_h=round(h*.16) if result.kind=='work_span' else round(h*.08) if result.educational_example else 0
    gx=x+caption_band;gy=y+legend_h;gw=w-caption_band;gh=h-footer_h-legend_h
    lookup={placed.id:placed for placed in result.nodes}
    bounds={}
    for placed in result.nodes:
        # A PowerPoint oval becomes distorted when normalized graph coordinates
        # are projected into a non-square visual region.  Re-center a true
        # circle in the allocated box and keep rounded semantic states wider.
        bounds[placed.id]=_graph_node_bounds(placed,gx,gy,gw,gh)
    def center(placed):
        nx,ny,nw,nh=bounds[placed.id];return nx+nw/2,ny+nh/2
    def endpoint(source,target,start):
        sx,sy=center(source);tx,ty=center(target)
        if start:
            px,py=_graph_boundary_point(bounds,source,tx,ty);return round(px),round(py)
        px,py=_graph_boundary_point(bounds,target,sx,sy)
        distance=max(1,((tx-sx)**2+(ty-sy)**2)**.5);gap=Pt(.5)
        # Stop shortly before the outline so the arrowhead does not merge with
        # the node border.  The marker still clearly points at the target.
        return round(px-(tx-sx)/distance*gap),round(py-(ty-sy)/distance*gap)
    if result.kind=='level_schedule':
        centers=_spread_positions(len(result.layer_captions),spec.rules.margin_y+spec.rules.node_height/2,1-spec.rules.margin_y-spec.rules.node_height/2)
        for index,center_y in enumerate(centers[:-1]):
            ly=gy+round(((center_y+centers[index+1])/2)*gh)
            separator=out.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,gx,ly,gx+gw,ly)
            separator.name=f'{node.id}-level-separator-{index+1}';separator.line.color.rgb=RGBColor.from_string(surface);separator.line.width=Pt(1)
    # Edges are drawn first so nodes remain visually dominant and hide tiny
    # endpoint inaccuracies introduced by PowerPoint's integer coordinates.
    for index,edge in enumerate(result.edges):
        source=lookup[edge.source];target=lookup[edge.target]
        x1,y1=endpoint(source,target,True);x2,y2=endpoint(source,target,False)
        if result.kind!='level_bound_proof':
            shape=out.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,x1,y1,x2,y2)
            shape.name=f'{node.id}-edge-{index+1}';shape.line.color.rgb=RGBColor.from_string(accent if edge.emphasis else dark)
            shape.line.width=Pt(2 if edge.emphasis else 1.25);_arrow(shape.line)
        if edge.label:
            mx=(x1+x2)//2;my=(y1+y2)//2;tw=Pt(28);th=Pt(18)
            label=out.shapes.add_textbox(mx-tw//2,my-th//2,tw,th);label.name=f'{node.id}-edge-label-{index+1}'
            tf=label.text_frame;tf.clear();tf.margin_left=tf.margin_right=0;tf.vertical_anchor=MSO_ANCHOR.MIDDLE
            p=tf.paragraphs[0];p.text=edge.label;p.alignment=PP_ALIGN.CENTER;p.font.name=node.style.font;p.font.size=Pt(12);p.font.bold=True;p.font.color.rgb=RGBColor.from_string(accent)
    for index,placed in enumerate(result.nodes):
        nx,ny,nw,nh=bounds[placed.id]
        fill=accent if placed.emphasis else surface;foreground=white if placed.emphasis else dark
        display=_math_display(placed.label if placed.weight is None else f'{placed.label}\n{placed.weight:g}')
        if placed.shape=='rounded':
            _card(out,node,nx,ny,nw,nh,display,f'{node.id}-graph-node-{placed.id}',fill,foreground,placed.emphasis)
        else:
            shape=out.shapes.add_shape(MSO_SHAPE.OVAL,nx,ny,nw,nh);shape.name=f'{node.id}-graph-node-{placed.id}'
            shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(fill);shape.line.color.rgb=RGBColor.from_string(accent);shape.line.width=Pt(1.4)
            _plain_shape(shape)
            if display:
                tf=shape.text_frame;tf.clear();tf.word_wrap=len(display)>3;tf.vertical_anchor=MSO_ANCHOR.MIDDLE
                tf.margin_left=tf.margin_right=Pt(2);tf.margin_top=tf.margin_bottom=Pt(1)
                p=tf.paragraphs[0];p.text=display;p.alignment=PP_ALIGN.CENTER;p.font.name=node.style.font;p.font.bold=True;p.font.size=Pt(max(10,min(16,_card_font_size(node,nw,nh,display,10))));p.font.color.rgb=RGBColor.from_string(foreground)
    if result.layer_captions:
        layer_count=len(result.layer_captions)
        centers=_spread_positions(layer_count,spec.rules.margin_y+spec.rules.node_height/2,1-spec.rules.margin_y-spec.rules.node_height/2)
        if result.direction=='BT':centers=list(reversed(centers))
        for index,(label,center_y) in enumerate(zip(result.layer_captions,centers)):
            tx=x;ty=gy+round((center_y-.08)*gh);tw=max(Pt(40),caption_band-round(w*.02));th=round(.16*gh)
            shape=out.shapes.add_textbox(tx,ty,tw,th);shape.name=f'{node.id}-layer-caption-{index+1}'
            tf=shape.text_frame;tf.clear();tf.word_wrap=True;tf.vertical_anchor=MSO_ANCHOR.MIDDLE;tf.margin_left=tf.margin_right=0
            p=tf.paragraphs[0];p.text=label;p.alignment=PP_ALIGN.LEFT;p.font.name=node.style.font
            p.font.size=Pt(min(13,_card_font_size(node,tw,th,label,10)));p.font.bold=True;p.font.color.rgb=RGBColor.from_string(accent)
    if result.footer:
        shape=out.shapes.add_textbox(x,y+h-footer_h,w,footer_h);shape.name=f'{node.id}-graph-footer'
        tf=shape.text_frame;tf.clear();tf.word_wrap=True;tf.vertical_anchor=MSO_ANCHOR.MIDDLE;tf.margin_left=tf.margin_right=Pt(2)
        footer=_math_display(result.footer)
        p=tf.paragraphs[0];p.text=footer;p.alignment=PP_ALIGN.CENTER;p.font.name=node.style.font
        p.font.size=Pt(min(14,_card_font_size(node,w,footer_h,footer,10)));p.font.bold=True;p.font.color.rgb=RGBColor.from_string(accent)
    if result.kind=='work_span':
        metrics=dict(result.metrics);labels=(f'Work W = {metrics.get("work",0):g}',f'Span S = {metrics.get("span",0):g}')
        for index,label in enumerate(labels):
            legend_gap=round(w*.03);legend_w=(w-legend_gap)//2
            lx=x+index*(legend_w+legend_gap)
            marker=out.shapes.add_shape(MSO_SHAPE.OVAL,lx,y+round(legend_h*.33),round(legend_h*.13),round(legend_h*.13))
            marker.name=f'{node.id}-legend-marker-{index+1}';marker.fill.solid();marker.fill.fore_color.rgb=RGBColor.from_string(accent if index else surface)
            marker.line.color.rgb=RGBColor.from_string(accent);marker.line.width=Pt(1.2);_plain_shape(marker)
            label_x=lx+round(legend_h*.18)
            shape=out.shapes.add_textbox(label_x,y,legend_w-round(legend_h*.18),legend_h);shape.name=f'{node.id}-legend-{index+1}'
            tf=shape.text_frame;tf.clear();tf.word_wrap=True;tf.vertical_anchor=MSO_ANCHOR.MIDDLE;tf.margin_left=tf.margin_right=Pt(2)
            p=tf.paragraphs[0];p.text=label;p.font.name=node.style.font
            p.font.size=Pt(min(11,_card_font_size(node,legend_w-round(legend_h*.18),legend_h,label,9)));p.font.bold=True;p.font.color.rgb=RGBColor.from_string(accent if index else dark)
    if result.educational_example:
        label_h=max(Pt(14),round(legend_h*.28));label=out.shapes.add_textbox(x,y,w,label_h);label.name=f'{node.id}-example-label'
        tf=label.text_frame;tf.clear();tf.margin_left=tf.margin_right=0
        p=tf.paragraphs[0];p.text=result.example_label or 'Учебный пример';p.alignment=PP_ALIGN.RIGHT;p.font.name=node.style.font;p.font.size=Pt(10);p.font.color.rgb=RGBColor.from_string(accent)
    return True


def _spread_positions(count,start,end):
    if count<=1:return [(start+end)/2]
    return [start+(end-start)*index/(count-1) for index in range(count)]


def _diagram(out,node,xywh):
    from .audit import contrast
    x,y,w,h=xywh;items=node.data.get('items',[])[:4];layout=node.data.get('layout','list')
    if not items:return
    accent=node.data.get('accent','0077FF');surface=node.data.get('surface') or node.style.fill or 'E8EEF6'
    dark='202020' if contrast('202020',surface)>=4.5 else 'FFFFFF'
    white='FFFFFF' if contrast('FFFFFF',accent)>=3 else '202020'
    gap=max(round(w*.025),Pt(8));vgap=max(round(h*.07),Pt(8))
    def connector(x1,y1,x2,y2,index):
        line=out.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,x1,y1,x2,y2)
        line.name=f'{node.id}-connector-{index}';line.line.color.rgb=RGBColor.from_string(accent);line.line.width=Pt(2)
    def dot(cx,cy,size,index,highlight=False):
        shape=out.shapes.add_shape(MSO_SHAPE.OVAL,cx-size//2,cy-size//2,size,size)
        shape.name=f'{node.id}-node-{index}';shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(accent if highlight else surface)
        shape.line.color.rgb=RGBColor.from_string(accent);shape.line.width=Pt(1.4)
    def flat_text(tx,ty,tw,th,value,name,color=dark,bold=False,align=PP_ALIGN.CENTER,size=None):
        shape=out.shapes.add_textbox(tx,ty,tw,th);shape.name=name
        tf=shape.text_frame;tf.clear();tf.word_wrap=True;tf.vertical_anchor=MSO_ANCHOR.MIDDLE
        tf.margin_left=tf.margin_right=Pt(3);tf.margin_top=tf.margin_bottom=Pt(2)
        p=tf.paragraphs[0];p.text=value;p.alignment=align;p.font.name=node.style.font
        fitted=_card_font_size(node,tw,th,value,minimum=12)
        p.font.size=Pt(min(fitted,size or fitted));p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
        return shape
    if _render_semantic_graph(out,node,xywh,layout,items,accent,surface,dark,white):return
    if layout=='pram_reality':
        # A semantic, fully editable comparison.  The labels are part of the
        # diagram grammar, so the renderer never asks a raster model to invent
        # hardware topology or technical text.
        half=round(w*.43);arrow_w=round(w*.08);header_h=round(h*.14)
        flat_text(x,y,half,header_h,'Идеальная PRAM',f'{node.id}-ideal-title',accent,True,size=19)
        flat_text(x+w-half,y,half,header_h,'Реальная система',f'{node.id}-real-title',accent,True,size=19)
        proc_y=y+round(h*.22);proc_h=round(h*.17);proc_gap=round(half*.025);proc_w=(half-2*proc_gap)//3
        for i in range(3):
            px=x+i*(proc_w+proc_gap)
            _card(out,node,px,proc_y,proc_w,proc_h,f'P{i+1}',f'{node.id}-ideal-p-{i+1}',surface,dark,True)
        memory_y=y+round(h*.63);memory_h=round(h*.22)
        for i in range(3):connector(x+i*(proc_w+proc_gap)+proc_w//2,proc_y+proc_h,x+half//2,memory_y,i+1)
        _card(out,node,x,memory_y,half,memory_h,'Общая память\nдоступ за один шаг',f'{node.id}-ideal-memory',accent,white,True)
        chevron=out.shapes.add_shape(MSO_SHAPE.CHEVRON,x+half+round(w*.015),y+round(h*.40),arrow_w,round(h*.18))
        chevron.name=f'{node.id}-reality-arrow';chevron.fill.solid();chevron.fill.fore_color.rgb=RGBColor.from_string(accent);chevron.line.fill.background()
        rx=x+w-half;cache_y=y+round(h*.36);cache_h=round(h*.14)
        for i in range(3):
            px=rx+i*(proc_w+proc_gap)
            _card(out,node,px,proc_y,proc_w,proc_h,f'P{i+1}',f'{node.id}-core-{i+1}',surface,dark,True)
            _card(out,node,px,cache_y,proc_w,cache_h,'Кэш',f'{node.id}-cache-{i+1}',surface,dark)
            connector(px+proc_w//2,proc_y+proc_h,px+proc_w//2,cache_y,i+10)
        bus_y=y+round(h*.62)
        for i in range(3):
            px=rx+i*(proc_w+proc_gap)
            connector(px+proc_w//2,cache_y+cache_h,px+proc_w//2,bus_y,i+30)
        connector(rx,bus_y,rx+half,bus_y,20)
        flat_text(rx,y+round(h*.51),half,round(h*.08),'Межсоединение · конкуренция',f'{node.id}-bus-label',accent,True,size=14)
        real_mem_y=y+round(h*.75)
        connector(rx+half//2,bus_y,rx+half//2,real_mem_y,21)
        _card(out,node,rx,real_mem_y,half,round(h*.16),'Память · задержки · overhead',f'{node.id}-real-memory',accent,white,True)
    elif layout=='abstraction':
        column_w=round(w*.38);arrow_w=round(w*.10);top_y=y+round(h*.08);label_h=round(h*.38)
        flat_text(x,top_y,column_w,label_h,'Concurrent-код\nСобытия и shared state',f'{node.id}-code',dark,True,size=19)
        flat_text(x+w-column_w,top_y,column_w,label_h,'Параллельный алгоритм\nWork и Span',f'{node.id}-model',dark,True,size=19)
        arrow=out.shapes.add_shape(MSO_SHAPE.CHEVRON,x+(w-arrow_w)//2,top_y+round(label_h*.28),arrow_w,round(label_h*.44))
        arrow.name=f'{node.id}-transition';arrow.fill.solid();arrow.fill.fore_color.rgb=RGBColor.from_string(accent);arrow.line.fill.background()
        metric='Полезный параллелизм W/S'
        line_y=y+round(h*.62)
        line=out.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,x+round(w*.12),line_y,x+round(w*.88),line_y)
        line.name=f'{node.id}-metric-line';line.line.color.rgb=RGBColor.from_string(accent);line.line.width=Pt(2)
        flat_text(x+round(w*.12),line_y+round(h*.04),round(w*.76),round(h*.24),metric,f'{node.id}-metric',accent,True,size=19)
    elif layout=='layers':
        count=len(items);ch=(h-vgap*(count-1))//count
        for i,item in enumerate(items):
            inset=round(i*w*.045);_card(out,node,x+inset,y+i*(ch+vgap),w-2*inset,ch,item,f'{node.id}-layer-{i+1}',accent if i==0 else surface,white if i==0 else dark,i==0)
    elif layout=='comparison':
        top_count=min(2,len(items));bottom=items[2:3]
        card_h=round(h*(.62 if bottom else .88));cw=(w-gap)//max(1,top_count)
        for i,item in enumerate(items[:top_count]):
            _card(out,node,x+i*(cw+gap),y,cw,card_h,item,f'{node.id}-compare-{i+1}',accent if i==0 else surface,white if i==0 else dark,i==0)
        if bottom:_card(out,node,x+round(w*.08),y+card_h+vgap,w-round(w*.16),h-card_h-vgap,bottom[0],f'{node.id}-compare-3',surface,dark)
    elif layout=='fork_join':
        size=max(Pt(24),min(round(w*.13),round(h*.16)));top_c=(x+w//2,y+round(h*.16));middle_y=y+round(h*.50);bottom_c=(x+w//2,y+round(h*.84))
        centers=[(x+round(w*.20),middle_y),(x+round(w*.50),middle_y),(x+round(w*.80),middle_y)]
        for i,(cx,cy) in enumerate(centers):connector(top_c[0],top_c[1]+size//2,cx,cy-size//2,i+1);connector(cx,cy+size//2,bottom_c[0],bottom_c[1]-size//2,i+10)
        dot(*top_c,size,0,True);dot(*bottom_c,size,9,True)
        for i,(cx,cy) in enumerate(centers):
            dot(cx,cy,size,i+1,False);flat_text(cx-size//2,cy-size//2,size,size,str(i+1),f'{node.id}-branch-label-{i+1}',dark,True,size=14)
        flat_text(x+round(w*.04),y,round(w*.30),round(h*.22),'Fork',f'{node.id}-fork-label',accent,True,size=20)
        flat_text(x+round(w*.04),y+round(h*.76),round(w*.30),round(h*.22),'Join',f'{node.id}-join-label',accent,True,size=20)
        flat_text(x+round(w*.62),y+round(h*.66),round(w*.34),round(h*.24),'Ready после предшественников',f'{node.id}-ready-label',dark,False,size=15)
    elif layout=='reduction_tree':
        tree_x=x+round(w*.40);tree_w=round(w*.58);levels=(4,2,1);ys=(y+round(h*.78),y+round(h*.44),y+round(h*.12));size=max(Pt(18),min(round(tree_w*.14),round(h*.13)))
        level_centers=[]
        for row,count in enumerate(levels):
            centers=[tree_x+round((i+1)*tree_w/(count+1)) for i in range(count)];level_centers.append(centers)
            for i,cx in enumerate(centers):dot(cx,ys[row],size,row*10+i,row==2)
        for child_row in (0,1):
            for i,cx in enumerate(level_centers[child_row]):
                parent=level_centers[child_row+1][i//2];connector(cx,ys[child_row]-size//2,parent,ys[child_row+1]+size//2,30+child_row*10+i)
        labels=('1. Копирование A в B','2. Попарное сложение','Сумма блока 2^h')
        label_h=round(h*.20);label_ys=(y+round(h*.69),y+round(h*.36),y+round(h*.03))
        for i,item in enumerate(labels):flat_text(x,label_ys[i],round(w*.34),label_h,item,f'{node.id}-level-label-{i+1}',accent if i==2 else dark,i==2,PP_ALIGN.LEFT,size=16)
    elif layout=='work_span':
        graph_w=round(w*.62);positions=[(.10,.76),(.35,.52),(.35,.84),(.60,.38),(.60,.68),(.88,.28),(.88,.58)]
        edges=((0,1),(0,2),(1,3),(1,4),(2,4),(3,5),(4,5),(4,6));critical={0,1,4,5};size=max(Pt(18),min(round(graph_w*.11),round(h*.12)))
        for i,(a,b) in enumerate(edges):connector(x+round(positions[a][0]*graph_w),y+round(positions[a][1]*h),x+round(positions[b][0]*graph_w),y+round(positions[b][1]*h),i+1)
        for i,(px,py) in enumerate(positions):dot(x+round(px*graph_w),y+round(py*h),size,i,i in critical)
        label_x=x+round(w*.66);label_w=round(w*.33)
        flat_text(label_x,y+round(h*.12),label_w,round(h*.24),items[0],f'{node.id}-work-label',dark,True,size=17)
        flat_text(label_x,y+round(h*.40),label_w,round(h*.24),items[1] if len(items)>1 else 'Span = критический путь',f'{node.id}-span-label',accent,True,size=17)
        if len(items)>2:flat_text(x+round(w*.08),y+round(h*.86),round(w*.84),round(h*.13),items[2],f'{node.id}-bound',accent,True,size=18)
    elif layout=='level_bound':
        row_ys=(y+round(h*.17),y+round(h*.43),y+round(h*.69));counts=(2,3,2);size=max(Pt(16),min(round(w*.10),round(h*.10)))
        for row,(cy,count) in enumerate(zip(row_ys,counts)):
            flat_text(x,cy-round(h*.08),round(w*.25),round(h*.16),f'Уровень {row+1}',f'{node.id}-level-{row+1}',dark,True,PP_ALIGN.LEFT,size=15)
            centers=[x+round(w*(.43+i*.19)) for i in range(count)]
            for i,cx in enumerate(centers):dot(cx,cy,size,row*10+i,row==0)
            if row:
                previous=[x+round(w*(.43+i*.19)) for i in range(counts[row-1])]
                for i,cx in enumerate(centers):connector(previous[min(i,len(previous)-1)],row_ys[row-1]+size//2,cx,cy-size//2,30+row*10+i)
        formula=items[2] if len(items)>2 else 'T_level ≤ 2T*'
        flat_text(x+round(w*.18),y+round(h*.82),round(w*.70),round(h*.17),formula,f'{node.id}-level-bound',accent,True,size=18)
    elif layout=='critical_path':
        positions=[(.10,.72),(.30,.50),(.50,.68),(.70,.35),(.90,.18)];size=max(Pt(18),min(round(w*.10),round(h*.14)))
        for i,((ax,ay),(bx,by)) in enumerate(zip(positions,positions[1:]),1):connector(x+round(ax*w),y+round(ay*h),x+round(bx*w),y+round(by*h),i)
        # A secondary branch makes the critical path visually explicit.
        connector(x+round(.30*w),y+round(.50*h),x+round(.52*w),y+round(.28*h),20);connector(x+round(.52*w),y+round(.28*h),x+round(.70*w),y+round(.35*h),21)
        dot(x+round(.52*w),y+round(.28*h),size,20,False)
        for i,(px,py) in enumerate(positions):dot(x+round(px*w),y+round(py*h),size,i,True)
        label_h=round(h*.27)
        for i,item in enumerate(items[:3]):
            if i==2:lx=x+round(.62*w);ly=y+round(.02*h);label_w=round(.36*w)
            else:lx=x+round((.01+i*.43)*w);ly=y+round(.72*h);label_w=round(.40*w)
            _card(out,node,lx,ly,label_w,label_h,item,f'{node.id}-path-label-{i+1}',accent if i==2 else surface,white if i==2 else dark,i==2)
    elif layout=='memory_access':
        processor_w=round(w*.30);memory_x=x+round(w*.54);memory_w=w-round(w*.56);card_h=round(h*.22)
        for i in range(3):
            py=y+i*round(h*.31);source_label=items[i] if i<len(items)-1 else ''
            modes=[mode for mode in ('EREW','CREW','CRCW') if mode.casefold() in source_label.casefold()]
            label='/'.join(modes) or 'Процессор'
            _card(out,node,x,py,processor_w,card_h,label,f'{node.id}-processor-{i+1}',accent if i==0 else surface,white if i==0 else dark,i==0)
            connector(x+processor_w,py+card_h//2,memory_x,y+h//2,i+1)
        memory_label=items[-1] if len(items)>1 else 'Общая память'
        _card(out,node,memory_x,y+round(h*.18),memory_w,round(h*.64),memory_label,f'{node.id}-memory',surface,dark,True)
    elif layout=='formula_focus':
        main_w=round(w*.52);main_h=round(h*.62);main_y=y+(h-main_h)//2
        _card(out,node,x,main_y,main_w,main_h,items[0],f'{node.id}-formula',accent,white,True)
        side_x=x+main_w+gap;side_w=w-main_w-gap;side_h=(h-vgap)//2
        for i,item in enumerate(items[1:3]):
            _card(out,node,side_x,y+i*(side_h+vgap),side_w,side_h,item,f'{node.id}-formula-note-{i+1}',surface,dark)
    elif layout=='scheduler':
        queue_w=round(w*.40);queue_h=round(h*.24)
        _card(out,node,x,y+round(h*.06),queue_w,queue_h,items[0],f'{node.id}-queue',accent,white,True)
        processor_x=x+round(w*.56);processor_w=w-round(w*.58);processor_h=round(h*.22)
        for i in range(3):
            py=y+i*round(h*.32);label=items[i+1] if i+1<len(items) else 'Процессор'
            connector(x+queue_w,y+round(h*.18),processor_x,py+processor_h//2,i+1)
            _card(out,node,processor_x,py,processor_w,processor_h,label,f'{node.id}-worker-{i+1}',surface,dark)
    elif layout=='sequence':
        count=len(items)
        if count>=3 and w/h<1.6:
            ch=(h-vgap*(count-1))//count;cx=x
            for i,item in enumerate(items):
                cy=y+i*(ch+vgap)
                if i:connector(x+w//2,cy-vgap,x+w//2,cy,i)
                _card(out,node,cx,cy,w,ch,item,f'{node.id}-card-{i+1}',accent if i==0 else surface,white if i==0 else dark,i==0)
        elif count>=4 or max(map(len,items),default=0)>52:
            cols=2;rows=(count+1)//2;cw=(w-gap)//2;ch=(h-vgap*(rows-1))//rows
            positions=[]
            for i,item in enumerate(items):
                row=i//2;col=i%2 if row%2==0 else 1-i%2
                positions.append((x+col*(cw+gap),y+row*(ch+vgap)))
            for i,((cx,cy),(nx,ny)) in enumerate(zip(positions,positions[1:]),1):
                connector(cx+cw//2,cy+ch//2,nx+cw//2,ny+ch//2,i)
            for i,(item,(cx,cy)) in enumerate(zip(items,positions)):
                _card(out,node,cx,cy,cw,ch,item,f'{node.id}-card-{i+1}',accent if i==0 else surface,white if i==0 else dark,i==0)
        else:
            cw=(w-gap*(count-1))//count;ch=min(round(h*.62),h);cy=y+(h-ch)//2
            for i,item in enumerate(items):
                cx=x+i*(cw+gap)
                if i:connector(cx-gap,cy+ch//2,cx,cy+ch//2,i)
                _card(out,node,cx,cy,cw,ch,item,f'{node.id}-card-{i+1}',accent if i==0 else surface,white if i==0 else dark,i==0)
    elif layout=='hierarchy':
        root_h=round(h*.30);root_w=round(w*.48);root_x=x+(w-root_w)//2
        child_count=max(1,len(items)-1);child_y=y+round(h*.57);child_h=h-(child_y-y)
        child_w=(w-gap*(child_count-1))//child_count
        for i in range(child_count):
            cx=x+i*(child_w+gap)
            connector(root_x+root_w//2,y+root_h,cx+child_w//2,child_y,i+1)
        _card(out,node,root_x,y,root_w,root_h,items[0],f'{node.id}-card-1',accent,white,True)
        for i,item in enumerate(items[1:] or items[:1]):
            cx=x+i*(child_w+gap)
            _card(out,node,cx,child_y,child_w,child_h,item,f'{node.id}-card-{i+2}',surface,dark)
    else:
        cols=1 if w/h<1.6 else 2 if len(items)>1 else 1;rows=(len(items)+cols-1)//cols
        cw=(w-gap*(cols-1))//cols;ch=(h-vgap*(rows-1))//rows
        for i,item in enumerate(items):
            col=i%cols;row=i//cols
            _card(out,node,x+col*(cw+gap),y+row*(ch+vgap),cw,ch,item,f'{node.id}-card-{i+1}',accent if i==0 else surface,white if i==0 else dark,i==0)


def temporary_objects(scene,slide):
    prs=Presentation();prs.slide_width=scene.width;prs.slide_height=scene.height
    out=prs.slides.add_slide(prs.slide_layouts[6])
    for node in slide.nodes:
        b=node.box;xywh=[round(v*s) for v,s in zip((b.x,b.y,b.w,b.h),(scene.width,scene.height,scene.width,scene.height))]
        if node.kind=='text':
            sh=out.shapes.add_textbox(*xywh);tf=sh.text_frame;tf.word_wrap=True
            if node.style.fill:
                sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(node.style.fill)
                tf.margin_left=tf.margin_right=Pt(10);tf.margin_top=tf.margin_bottom=Pt(8)
            else:tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
            for i,line in enumerate(_math_display(node.text).split('\n')):
                p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
                p.text=line;p.font.name=node.style.font;p.font.size=Pt(node.style.size);p.font.bold=node.style.bold;p.font.color.rgb=RGBColor.from_string(node.style.color)
                p.alignment={'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}.get(node.style.align,PP_ALIGN.LEFT)
        elif node.kind=='chart':
            data=CategoryChartData();data.categories=node.data['categories']
            for name,values in node.data['series'].items():data.add_series(name,values)
            typ=XL_CHART_TYPE.LINE_MARKERS if scene.variant=='B' else XL_CHART_TYPE.COLUMN_CLUSTERED
            sh=out.shapes.add_chart(typ,*xywh,data);chart=sh.chart
            chart.has_legend=len(node.data['series'])>1
            if chart.has_legend:chart.legend.position=XL_LEGEND_POSITION.BOTTOM
            chart.value_axis.has_title=True;chart.value_axis.axis_title.text_frame.text=node.data['unit']
            chart.category_axis.has_title=True;chart.category_axis.axis_title.text_frame.text=node.data['title']
            for axis in (chart.category_axis,chart.value_axis):
                axis.tick_labels.font.name=node.style.font;axis.tick_labels.font.size=Pt(10);axis.tick_labels.font.color.rgb=RGBColor.from_string(node.style.color)
                for para in axis.axis_title.text_frame.paragraphs:
                    para.font.name=node.style.font;para.font.size=Pt(12);para.font.color.rgb=RGBColor.from_string(node.style.color)
            if chart.has_legend:
                chart.legend.font.name=node.style.font;chart.legend.font.size=Pt(10);chart.legend.font.color.rgb=RGBColor.from_string(node.style.color)
            for i,series in enumerate(chart.series):
                if typ==XL_CHART_TYPE.COLUMN_CLUSTERED:
                    series.format.fill.solid();series.format.fill.fore_color.rgb=RGBColor.from_string(node.style.color)
        elif node.kind=='table':
            cols=1+len(node.data['series']);rows=1+len(node.data['categories'])
            sh=out.shapes.add_table(rows,cols,*xywh)
            headers=[node.data['title']]+[name+' ('+node.data['unit']+')' for name in node.data['series']]
            values=[headers]+[[cat]+[str(v[i]) for v in node.data['series'].values()] for i,cat in enumerate(node.data['categories'])]
            for i,row in enumerate(values):
                for j,value in enumerate(row):
                    cell=sh.table.cell(i,j);cell.text=value
                    for p in cell.text_frame.paragraphs:
                        p.font.name=node.style.font;p.font.size=Pt(max(8,min(node.style.size,14)));p.font.color.rgb=RGBColor.from_string(node.style.color)
        elif node.kind=='image':
            source=artifact_path(node.data['path'])
            from PIL import Image,ImageChops
            picture=str(source)
            with Image.open(source) as opened:
                im=opened.convert('RGB')
                if str(node.data.get('source','')).startswith('Z-Image'):
                    background=Image.new('RGB',im.size,im.getpixel((0,0)))
                    diff=ImageChops.difference(im,background).convert('L').point(lambda value:255 if value>18 else 0)
                    bbox=diff.getbbox()
                    if bbox:
                        pad=max(12,round(min(im.size)*.04));l,t,r,bt=bbox
                        im=im.crop((max(0,l-pad),max(0,t-pad),min(im.width,r+pad),min(im.height,bt+pad)))
                    stream=io.BytesIO();im.save(stream,format='PNG');stream.seek(0);picture=stream
                iw,ih=im.size
            factor=min(xywh[2]/iw,xywh[3]/ih)
            w,h=round(iw*factor),round(ih*factor)
            sh=out.shapes.add_picture(picture,xywh[0]+(xywh[2]-w)//2,xywh[1]+(xywh[3]-h)//2,w,h)
        elif node.kind=='diagram':
            _diagram(out,node,xywh)
            continue
        else:continue
        sh.name=node.id
    buf=io.BytesIO();prs.save(buf);return Package(buf.getvalue())


def strip_content(root,prototype):
    remove={s.shape_id for s in prototype.slots if s.role in {'title','body','visual'}}
    # Sample decks often contain phones, screenshots and empty demo frames built
    # from ordinary vector shapes. Keep only small edge decorations and footers;
    # the master/layout still carries the brand background, logo and page chrome.
    for slot in prototype.slots:
        area=slot.box.w*slot.box.h
        edge=slot.box.x<.035 or slot.box.y<.035 or slot.box.x+slot.box.w>.965 or slot.box.y+slot.box.h>.965
        if slot.role=='decor' and not (edge and area<=.02):remove.add(slot.shape_id)
    for shape in list(root.findall('.//p:sp',NS))+list(root.findall('.//p:pic',NS))+list(root.findall('.//p:graphicFrame',NS)):
        nv=shape.find('.//p:cNvPr',NS)
        if nv is not None and int(nv.get('id')) in remove:shape.getparent().remove(shape)


def export_pptx(template: bytes,design: DesignIR,scene: SceneIR,path: Path,decor_only=False):
    source=Package(template);pkg=Package(template);prototypes={p.id:p for p in design.prototypes};slideparts=[]
    for i,slide in enumerate(scene.slides):
        proto=prototypes[slide.prototype_id]
        dest=pkg.clone_graph(source,proto.slide_part,f'v{scene.version}_{scene.variant}_{i}')
        root=pkg.root(dest);strip_content(root,proto)
        tree=root.find('p:cSld/p:spTree',NS)
        # Speaker notes from a template may contain unrelated content.
        rr=pkg.root(relpath(dest))
        for r in list(rr):
            if r.get('Type','').endswith('/notesSlide'):rr.remove(r)
        pkg.parts[relpath(dest)]=serialize(rr)
        if not decor_only:
            objects=temporary_objects(scene,slide)
            temp_part=objects.slides()[0]
            objroot=objects.root(temp_part)
            rr=pkg.root(relpath(dest));nextid=max([int(n.get('id','0')) for n in tree.findall('.//p:cNvPr',NS)]+[1])+1
            for shape in objroot.find('p:cSld/p:spTree',NS):
                if shape.tag.split('}')[-1] not in {'sp','cxnSp','pic','graphicFrame'}:continue
                shape=copy.deepcopy(shape)
                for nv in shape.findall('.//p:cNvPr',NS):nv.set('id',str(nextid));nextid+=1
                for el in shape.iter():
                    for attr,value in list(el.attrib.items()):
                        if attr.startswith('{'+NS['r']+'}'):
                            target=objects.target(temp_part,value)
                            rel=next(r for r in objects.relations(temp_part) if r.get('Id')==value)
                            newrid=f'rIdObject{i}_{nextid}_{value}'
                            newrel=copy.deepcopy(rel);newrel.set('Id',newrid)
                            if target:
                                newtarget=pkg.clone_graph(objects,target,f'v{scene.version}_{scene.variant}_{i}_{nextid}')
                                newrel.set('Target',posixpath.relpath(newtarget,posixpath.dirname(dest)))
                            rr.append(newrel);el.set(attr,newrid)
                tree.append(shape)
            pkg.parts[relpath(dest)]=serialize(rr)
            for node in slide.nodes:
                if node.kind=='smartart':
                    tree.append(inject(pkg,dest,node,nextid,scene.width,scene.height));nextid+=1
        pkg.parts[dest]=serialize(root);slideparts.append(dest)
    pkg.set_slides(slideparts);pkg.prune();pkg.validate();path.parent.mkdir(parents=True,exist_ok=True);path.write_bytes(pkg.bytes())
    # Independent native object reader.
    assert len(Presentation(str(path)).slides)==len(scene.slides)


def render(pptx: Path,outdir: Path,timeout=90):
    outdir.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='vktech-lo-') as profile:
        command=[os.environ.get('SOFFICE','soffice'),'-env:UserInstallation='+Path(profile).as_uri(),'--headless','--convert-to','pdf','--outdir',str(outdir),str(pptx)]
        result=subprocess.run(command,capture_output=True,timeout=timeout)
        pdf=outdir/(pptx.stem+'.pdf')
        if result.returncode or not pdf.exists():raise RuntimeError('LibreOffice export failed: '+result.stderr.decode(errors='replace')[-400:])
    from pypdf import PdfReader
    pages=len(PdfReader(pdf).pages)
    result=subprocess.run([os.environ.get('PDFTOPPM','pdftoppm'),'-scale-to','1400','-png',str(pdf),str(outdir/'slide')],capture_output=True,timeout=timeout)
    if result.returncode:raise RuntimeError('PDF preview rendering failed')
    images=sorted(outdir.glob('slide-*.png'),key=lambda p:int(p.stem.rsplit('-',1)[-1]))
    if len(images)!=pages:raise RuntimeError('Preview page count mismatch')
    return pdf,images


def chart_svg(node,variant):
    data=node.data;cats=data['categories'];series=list(data['series'].items());values=[v for _,vs in series for v in vs]
    low=min(0,min(values));high=max(0,max(values));span=high-low or 1
    y=lambda v:230-(v-low)/span*190
    parts=[f'<svg viewBox="0 0 600 280" role="img" aria-label="{html.escape(data["title"],quote=True)}"><title>{html.escape(data["title"])} ({html.escape(data["unit"])})</title>',f'<line x1="50" y1="{y(0)}" x2="580" y2="{y(0)}" stroke="currentColor"/>']
    for si,(name,vs) in enumerate(series):
        points=[]
        for i,v in enumerate(vs):
            x=60+(i+.5)*510/len(cats)
            if variant=='B':points.append(f'{x},{y(v)}')
            else:
                bw=450/len(cats)/len(series);xx=x-225/len(cats)+si*bw
                parts.append(f'<rect x="{xx}" y="{min(y(v),y(0))}" width="{bw*.85}" height="{abs(y(v)-y(0))}" fill="#{node.style.color}" opacity="{1-si/max(1,len(series))*0.5}"><title>{html.escape(name)}: {v}</title></rect>')
            parts.append(f'<text x="{x}" y="{y(v)-5}" font-size="11" text-anchor="middle">{v}</text>')
        if variant=='B':parts.append(f'<polyline points="{" ".join(points)}" fill="none" stroke="#{node.style.color}" stroke-width="3"/>')
    for i,cat in enumerate(cats):parts.append(f'<text x="{60+(i+.5)*510/len(cats)}" y="255" font-size="11" text-anchor="middle">{html.escape(cat)}</text>')
    parts.append(f'<text x="50" y="20" font-size="12">{html.escape(data["unit"])}</text></svg>')
    return ''.join(parts)


def graph_svg(node):
    """Render the same semantic layout used by PPTX as accessible SVG."""
    from .audit import contrast
    layout_name=node.data.get('layout','list');items=node.data.get('items',[])
    spec=semantic_graph(layout_name,items,node.data.get('graph'))
    if spec is None:return ''
    result=layout_graph(spec);accent=node.data.get('accent','0077FF');surface=node.data.get('surface') or node.style.fill or 'E8EEF6'
    dark='202020' if contrast('202020',surface)>=4.5 else 'FFFFFF';on_accent='FFFFFF' if contrast('FFFFFF',accent)>=3 else '202020'
    caption=92 if result.layer_captions else 10;top=36 if result.kind=='work_span' else 24 if result.educational_example else 10;footer=28 if result.footer else 8
    gx,gy,gw,gh=caption,top,600-caption-15,300-top-footer
    bounds={}
    for placed in result.nodes:
        x=gx+placed.x*gw;y=gy+placed.y*gh;w=placed.w*gw;h=placed.h*gh
        if placed.shape=='circle':
            d=min(w,h);x+=(w-d)/2;y+=(h-d)/2;w=h=d
        bounds[placed.id]=(x,y,w,h)
    def center(placed):
        x,y,w,h=bounds[placed.id];return x+w/2,y+h/2
    def boundary(placed,toward):
        x,y,w,h=bounds[placed.id];cx,cy=x+w/2,y+h/2;dx,dy=toward[0]-cx,toward[1]-cy
        if not dx and not dy:return cx,cy
        scale=(w/2)/max(1,(dx*dx+dy*dy)**.5) if placed.shape=='circle' else 1/max(abs(dx)/(w/2),abs(dy)/(h/2))
        return cx+dx*scale,cy+dy*scale
    lookup={placed.id:placed for placed in result.nodes}
    title=html.escape(node.data.get('description') or layout_name)
    marker_id=html.escape(node.id,quote=True)
    parts=[f'<svg viewBox="0 0 600 300" role="img" aria-label="{title}"><title>{title}</title>',f'<defs><marker id="arrow-dark-{marker_id}" viewBox="0 0 10 10" refX="9" refY="5" markerWidth="5" markerHeight="5" orient="auto-start-reverse"><path d="M 0 0 L 10 5 L 0 10 z" fill="#{dark}"/></marker><marker id="arrow-accent-{marker_id}" viewBox="0 0 10 10" refX="9" refY="5" markerWidth="5" markerHeight="5" orient="auto-start-reverse"><path d="M 0 0 L 10 5 L 0 10 z" fill="#{accent}"/></marker></defs>']
    if result.kind=='level_schedule' and len(result.layer_captions)>1:
        centers=_spread_positions(len(result.layer_captions),spec.rules.margin_y+spec.rules.node_height/2,1-spec.rules.margin_y-spec.rules.node_height/2)
        for index in range(len(centers)-1):
            y=gy+(centers[index]+centers[index+1])/2*gh;parts.append(f'<line x1="{gx}" y1="{y:.1f}" x2="{gx+gw}" y2="{y:.1f}" stroke="#{surface}"/>')
    for edge in result.edges:
        source=lookup[edge.source];target=lookup[edge.target];sc=center(source);tc=center(target);x1,y1=boundary(source,tc);x2,y2=boundary(target,sc)
        color=accent if edge.emphasis else dark;width=3 if edge.emphasis else 1.7
        if result.kind!='level_bound_proof':parts.append(f'<line x1="{x1:.1f}" y1="{y1:.1f}" x2="{x2:.1f}" y2="{y2:.1f}" stroke="#{color}" stroke-width="{width}" marker-end="url(#arrow-{"accent" if edge.emphasis else "dark"}-{marker_id})"/>')
        if edge.label:parts.append(f'<text x="{(x1+x2)/2:.1f}" y="{(y1+y2)/2-5:.1f}" font-size="13" font-weight="700" text-anchor="middle" fill="#{accent}">{html.escape(edge.label)}</text>')
    for placed in result.nodes:
        x,y,w,h=bounds[placed.id];fill=accent if placed.emphasis else surface;color=on_accent if placed.emphasis else dark
        if placed.shape=='circle':parts.append(f'<circle cx="{x+w/2:.1f}" cy="{y+h/2:.1f}" r="{w/2:.1f}" fill="#{fill}" stroke="#{accent}" stroke-width="1.5"/>')
        else:parts.append(f'<rect x="{x:.1f}" y="{y:.1f}" width="{w:.1f}" height="{h:.1f}" rx="10" fill="#{fill}" stroke="#{accent}" stroke-width="1.5"/>')
        lines=[_math_display(placed.label)]+([f'{placed.weight:g}'] if placed.weight is not None else []);base=y+h/2-(len(lines)-1)*7
        for index,value in enumerate(lines):parts.append(f'<text x="{x+w/2:.1f}" y="{base+index*15:.1f}" font-size="13" font-weight="700" text-anchor="middle" dominant-baseline="middle" fill="#{color}">{html.escape(value)}</text>')
    if result.layer_captions:
        centers=_spread_positions(len(result.layer_captions),spec.rules.margin_y+spec.rules.node_height/2,1-spec.rules.margin_y-spec.rules.node_height/2)
        if result.direction=='BT':centers=list(reversed(centers))
        for label,center_y in zip(result.layer_captions,centers):parts.append(f'<text x="8" y="{gy+center_y*gh:.1f}" font-size="13" font-weight="700" dominant-baseline="middle" fill="#{accent}">{html.escape(label)}</text>')
    if result.kind=='work_span':
        metrics=dict(result.metrics);parts.append(f'<text x="15" y="20" font-size="13" font-weight="700" fill="#{dark}">Work W = {metrics.get("work",0):g}</text><text x="210" y="20" font-size="13" font-weight="700" fill="#{accent}">Span S = {metrics.get("span",0):g}</text>')
    if result.educational_example:parts.append(f'<text x="585" y="20" font-size="11" text-anchor="end" fill="#{accent}">{html.escape(result.example_label or "Учебный пример")}</text>')
    if result.footer:parts.append(f'<text x="300" y="292" font-size="13" font-weight="700" text-anchor="middle" fill="#{accent}">{html.escape(_math_display(result.footer))}</text>')
    parts.append('</svg>');return ''.join(parts)


def export_html(scene,path,backgrounds=None):
    slides=[]
    for i,slide in enumerate(scene.slides):
        bg=''
        if backgrounds:
            bg='<img class="decor" aria-hidden="true" alt="" src="data:image/png;base64,'+base64.b64encode(backgrounds[i].read_bytes()).decode()+'">'
        nodes=[]
        for n in slide.nodes:
            b=n.box;fs=n.style.size/72/(scene.width/914400)*100
            style=f'left:{b.x*100}%;top:{b.y*100}%;width:{b.w*100}%;height:{b.h*100}%;font-family:{html.escape(json.dumps(n.style.font),quote=True)},sans-serif;font-size:{fs}cqw;color:#{n.style.color};text-align:{n.style.align};font-weight:{700 if n.style.bold else 400};background:{"#"+n.style.fill if n.style.fill else "transparent"};padding:{"0.45em" if n.style.fill else "0"}'
            if n.kind=='text':body=html.escape(_math_display(n.text)).replace('\n','<br>')
            elif n.kind=='chart':body=chart_svg(n,scene.variant)
            elif n.kind=='table':
                data=n.data;headers=[data['title']]+[name+' ('+data['unit']+')' for name in data['series']]
                rows=[headers]+[[cat]+[str(v[j]) for v in data['series'].values()] for j,cat in enumerate(data['categories'])]
                body='<table>'+''.join('<tr>'+''.join(f'<{"th" if ri==0 else "td"}>{html.escape(cell)}</{"th" if ri==0 else "td"}>' for cell in row)+'</tr>' for ri,row in enumerate(rows))+'</table>'
            elif n.kind=='diagram' and semantic_graph(n.data.get('layout','list'),n.data.get('items',[]),n.data.get('graph')):
                body=graph_svg(n)
            elif n.kind in {'diagram','smartart'}:
                layout=n.data.get('layout','list');orientation=' vertical' if layout=='sequence' and n.box.w/n.box.h<1.6 else ''
                body='<ol class="diagram '+layout+orientation+'">'+''.join('<li>'+html.escape(t)+'</li>' for t in n.data.get('items',[]))+'</ol>'
            elif n.kind=='image':body='<img alt="'+html.escape(n.data['description'],quote=True)+'" src="data:'+n.data.get('media_type','image/png')+';base64,'+base64.b64encode(artifact_path(n.data['path']).read_bytes()).decode()+'">'
            else:body=''
            nodes.append(f'<div class="node {n.kind}" data-element-id="{html.escape(n.id,quote=True)}" style="{style}">{body}</div>')
        slides.append(f'<section class="slide" aria-label="{html.escape(slide.title,quote=True)}" style="background:#{slide.background}">{bg}{"".join(nodes)}</section>')
    doc='''<!doctype html><html lang="ru"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Презентация</title><style>
body{margin:0;background:#d9dce3;font-family:Arial}main{max-width:1400px;margin:auto}.slide{position:relative;aspect-ratio:16/9;container-type:inline-size;margin:24px 0;overflow:hidden}.decor{position:absolute;width:100%;height:100%;inset:0}.node{position:absolute;box-sizing:border-box;line-height:1.2;overflow:visible}.node img,.node svg{width:100%;height:100%;object-fit:contain}.node table{border-collapse:collapse;width:100%;height:100%;font-size:.7em}.node td,.node th{padding:.2em;border:1px solid currentColor}.diagram{list-style:none;padding:0;margin:0;display:flex;flex-direction:column;height:100%;gap:4%;font-size:.8em}.diagram li{background:#e8eef6;padding:.4em;border-radius:8px;flex:1}.diagram.sequence{flex-direction:row}.diagram.sequence.vertical,.diagram.layers{flex-direction:column}.diagram.comparison{display:grid;grid-template-columns:1fr 1fr;grid-auto-rows:1fr}.diagram.fork_join,.diagram.reduction_tree,.diagram.critical_path{justify-content:space-between;align-items:center}.diagram.fork_join li,.diagram.reduction_tree li,.diagram.critical_path li{width:70%}.diagram.memory_access,.diagram.scheduler,.diagram.formula_focus{display:grid;grid-template-columns:1fr 1fr}.controls{position:fixed;right:20px;bottom:20px;z-index:20;background:#172438;color:white;border:0;border-radius:8px;padding:10px 14px}.controls button{color:white;background:transparent;border:0;font-size:18px}.controls span{padding:0 8px}@media print{body{background:white}.slide{margin:0;break-after:page}.controls{display:none}main{max-width:none}}@supports not (font-size:1cqw){.node{font-size:18px!important}}
</style><main>'''+''.join(slides)+'''</main><nav class="controls" aria-label="Навигация"><button type="button" data-dir="-1" aria-label="Предыдущий слайд">←</button><span></span><button type="button" data-dir="1" aria-label="Следующий слайд">→</button></nav><script>
const slides=[...document.querySelectorAll('.slide')],label=document.querySelector('.controls span');let index=0;
function show(i){index=Math.max(0,Math.min(slides.length-1,i));slides.forEach((s,j)=>s.hidden=j!==index);label.textContent=`${index+1} / ${slides.length}`;slides[index].focus?.()}
document.querySelectorAll('[data-dir]').forEach(b=>b.onclick=()=>show(index+Number(b.dataset.dir)));addEventListener('keydown',e=>{if(e.key==='ArrowRight'||e.key==='PageDown')show(index+1);if(e.key==='ArrowLeft'||e.key==='PageUp')show(index-1)});show(0);
</script></html>'''
    path.write_text(doc,encoding='utf-8')
